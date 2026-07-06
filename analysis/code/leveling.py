"""Core data loading and analysis for the leveling-effect study (NS372 and delta22). No plotting.

Reads the two datasets, decodes the stored shieldings, and runs the per-nucleus correlation / PCA
analysis. The figure notebook (built by build_nb_leveling.py) and test_leveling.py import from here.

The claim tested, on each dataset separately (never pooled): once a linear scaling is allowed for,
any two shielding methods are near-perfectly correlated, PC1 carries almost all between-method
variance, and the methods fall on a 1D curve in (PC1, PC2).
"""
import numpy as np
import pandas as pd
import h5py
import openpyxl
from sklearn.decomposition import PCA
from sklearn.preprocessing import StandardScaler

from fixed_point import decode_fixed_point as _decode_fixed_point
from stats import rmse as _rmse

# canonical functional-family ordering used to sort method columns in every figure
FAMILY_ORDER = ['LDA', 'GGA', 'mGGA', 'GH', 'RSH', 'LH', 'DH', 'WFT', 'ref']

def sort_by_family(M, methods, families):
    """Stable-sort method columns of M into FAMILY_ORDER (reference family last)."""
    rank = {f: i for i, f in enumerate(FAMILY_ORDER)}
    idx = sorted(range(len(methods)), key=lambda i: (rank.get(families[i], 99), i))
    return M[:, idx], [methods[i] for i in idx], [families[i] for i in idx]

def log_r(R):
    """-log10(1-|r|): 1 -> r=0.9, 2 -> 0.99, 3 -> 0.999, ..."""
    return -np.log10(np.clip(1 - np.abs(R), 1e-6, None))

def scaled_rmse(method, ref):
    """RMSE of residuals after the per-method linear fit method ~ a*ref + b."""
    a, b = np.polyfit(ref, method, 1)
    return _rmse(method, a * ref + b)

def _parabola_fit(xn, yn, n_angles=720):
    # best degree-2 fit  y_rot ~ x_rot  searched over rotation angle theta
    best = (0.0, -np.inf, None)
    for th in np.linspace(0, np.pi, n_angles):
        c, s = np.cos(th), np.sin(th)
        xr, yr = c * xn + s * yn, -s * xn + c * yn
        A = np.column_stack([np.ones_like(xr), xr, xr ** 2])
        coef, *_ = np.linalg.lstsq(A, yr, rcond=None)
        ss_res = ((yr - A @ coef) ** 2).sum()
        ss_tot = ((yr - yr.mean()) ** 2).sum()
        r2 = 1 - ss_res / ss_tot if ss_tot > 0 else 0.0
        if r2 > best[1]:
            best = (th, r2, coef)
    return best

def analyze_nucleus(M, methods, ref=None):
    """Correlation, PCA, and parabola-fit diagnostics for one nucleus.

    M is an (n_observations, n_methods) shielding matrix; ref (optional) is the reference method's
    values, used to compute each method's scaled RMSE against it.
    """
    R = np.corrcoef(M, rowvar=False)
    iu = np.triu_indices_from(R, k=1)
    rv = R[iu]
    Z = StandardScaler().fit_transform(M)
    pca = PCA().fit(Z)
    ev = pca.explained_variance_ratio_
    pc1, pc2 = pca.components_[0].copy(), pca.components_[1].copy()
    if pc1.mean() < 0:                       # cosmetic sign convention
        pc1 = -pc1
    mx, my = pc1.mean(), pc2.mean()
    sx, sy = max(pc1.std(), 1e-12), max(pc2.std(), 1e-12)
    xn, yn = (pc1 - mx) / sx, (pc2 - my) / sy
    th, pr2, coef = _parabola_fit(xn, yn)
    c, s = np.cos(th), np.sin(th)
    par = dict(theta=th, coef=coef, r2=pr2, mx=mx, sx=sx, my=my, sy=sy,
               xr=c * xn + s * yn)
    out = dict(R=R, methods=list(methods), pc1=pc1, pc2=pc2, ev=ev, par=par,
               r_min=float(rv.min()), r_med=float(np.median(rv)),
               r_max=float(rv.max()))
    if ref is not None:
        out['scaled_rmse'] = {m: scaled_rmse(M[:, i], ref)
                              for i, m in enumerate(methods)}
    return out

def parabola_curve(par, npts=300):
    """Sample points along the fitted parabola (from analyze_nucleus's `par`) in PC1/PC2 space."""
    th, coef = par['theta'], par['coef']
    c, s = np.cos(th), np.sin(th)
    xr = par['xr']
    vertex = -coef[1] / (2 * coef[2]) if coef[2] != 0 else float(xr.mean())
    lo, hi = min(xr.min(), vertex), max(xr.max(), vertex)
    pad = 0.15 * (hi - lo + 1e-9)
    t = np.linspace(lo - pad, hi + pad, npts)
    yr = coef[0] + coef[1] * t + coef[2] * t ** 2
    xn, yn = c * t - s * yr, s * t + c * yr
    return xn * par['sx'] + par['mx'], yn * par['sy'] + par['my']

def summarize(data, res):
    """Per-nucleus leveling diagnostics (analyze_nucleus results) as one summary DataFrame."""
    import pandas as pd
    rows = []
    for nuc, r in res.items():
        rows.append(dict(nucleus=nuc,
                         n_obs=data[nuc]['M'].shape[0],
                         n_methods=data[nuc]['M'].shape[1],
                         r_min=r['r_min'], r_median=r['r_med'],
                         PC1_pct=r['ev'][0] * 100, PC2_pct=r['ev'][1] * 100,
                         PC3plus_pct=(1 - r['ev'][0] - r['ev'][1]) * 100,
                         parabola_R2=r['par']['r2']))
    return pd.DataFrame(rows)

def dataset_logr_max(res):
    """Largest -log10(1-|r|) over every nucleus of a dataset (lower triangle of each R)."""
    vm = 0.0
    for r in res.values():
        nm = len(r['methods']); m = nm - 1
        sub = r['R'][1:, :-1]
        hide = np.triu(np.ones((m, m), dtype=bool), k=1)
        vm = max(vm, float(np.nanmax(np.where(hide, np.nan, log_r(sub)))))
    return float(np.ceil(vm * 10) / 10)


# ----------------------------------------------------------------------------
# NS372 (Kaupp) loader
# ----------------------------------------------------------------------------
# (functional, tau-treatment, %HF, family, display label)
# one tau-variant per functional: tauC (Gaussian default) for tau-dependent
# meta-GGAs / hybrids, tauD for local hybrids. Match keys use real Unicode.
METHODS = [
    ('SVWN', None, 0, 'LDA', 'SVWN'),
    ('BLYP', None, 0, 'GGA', 'BLYP'), ('BP86', None, 0, 'GGA', 'BP86'),
    ('PBE', None, 0, 'GGA', 'PBE'), ('B97D', None, 0, 'GGA', 'B97D'),
    ('HCTH', None, 0, 'GGA', 'HCTH'),
    ('KT1', None, 0, 'GGA', 'KT1'), ('KT2', None, 0, 'GGA', 'KT2'), ('KT3', None, 0, 'GGA', 'KT3'),
    ('TPSS', 'τC', 0, 'mGGA', 'TPSS/τC'), ('τ-HCTH', 'τC', 0, 'mGGA', 'τ-HCTH/τC'),
    ('M06-L', 'τC', 0, 'mGGA', 'M06-L/τC'), ('VSXC', 'τC', 0, 'mGGA', 'VSXC/τC'),
    ('MN15-L', 'τC', 0, 'mGGA', 'MN15-L/τC'), ('B97M-V', 'τC', 0, 'mGGA', 'B97M-V/τC'),
    ('SCAN', 'τC', 0, 'mGGA', 'SCAN/τC'), ('rSCAN', 'τC', 0, 'mGGA', 'rSCAN/τC'),
    ('r2SCAN', 'τC', 0, 'mGGA', 'r2SCAN/τC'),
    ('TPSSh', 'τC', 10, 'GH', 'TPSSh/τC'), ('B3LYP', None, 20, 'GH', 'B3LYP'),
    ('B97-2', None, 21, 'GH', 'B97-2'), ('PBE0', None, 25, 'GH', 'PBE0'),
    ('M06', 'τC', 27, 'GH', 'M06/τC'), ('PW6B95', 'τC', 28, 'GH', 'PW6B95/τC'),
    ('MN15', 'τC', 44, 'GH', 'MN15/τC'), ('BHLYP', None, 50, 'GH', 'BHLYP'),
    ('M06-2X', 'τC', 54, 'GH', 'M06-2X/τC'),
    ('CAM-B3LYP', None, 65, 'RSH', 'CAM-B3LYP'),
    ('ωB97X-D', None, 100, 'RSH', 'ωB97X-D'), ('ωB97X-V', None, 100, 'RSH', 'ωB97X-V'),
    ('ωB97M-V', 'τC', 100, 'RSH', 'ωB97M-V/τC'),
    ('LH07s-SVWN', None, 22, 'LH', 'LH07s-SVWN'),
    ('MPSTS', 'τD', 30, 'LH', 'mPSTS/τD'),
    ('LHJ14', 'τD', 10, 'LH', 'LHJ14/τD'),
    ('LH07t-SVWN', 'τD', 48, 'LH', 'LH07t-SVWN/τD'),
    ('LH12ct-SsirPW92', 'τD', 65, 'LH', 'LH12ct-SsirPW92/τD'),
    ('LH12ct-SsifPW92', 'τD', 71, 'LH', 'LH12ct-SsifPW92/τD'),
    ('LH14t-calPBE', 'τD', 50, 'LH', 'LH14t-calPBE/τD'),
    ('LH20t', 'τD', 71, 'LH', 'LH20t/τD'),
    ('B2PLYP', None, 53, 'DH', 'B2PLYP'), ('B2GP-PLYP', None, 65, 'DH', 'B2GP-PLYP'),
    ('DSD-PBEP86', None, 70, 'DH', 'DSD-PBEP86'),
    ('Hartree-Fock', None, 100, 'WFT', 'HF'), ('MP2', None, 100, 'WFT', 'MP2'),
]
# (display nucleus, source sheet, in-sheet nucleus filter, molecules to drop)
NUCLEI = [
    ('1H',  'S6 - 1H Shieldings',  None, ()),
    ('11B', 'S7 - 11B Shieldings', None, ('BH',)),
    ('13C', 'S8 - 13C Shieldings', None, ()),
    ('15N', 'S9 - 15N Shieldings', None, ()),
    ('17O', 'S10 - 17O Shieldings', None, ('O3',)),
    ('19F', 'S11 - 19F Shieldings', None, ('F3-',)),
    ('31P', 'S12 - 31P - 33S Shieldings', '31P', ()),
    ('33S', 'S12 - 31P - 33S Shieldings', '33S', ()),
]

def read_nucleus(wb, sheet, nuc_filter=None, exclude=()):
    """Read the CCSD(T) reference plus each method's shielding vector from one Kaupp SI sheet."""
    ws = wb[sheet]
    rows = []; last_nuc = None
    for r in range(4, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v:
            last_nuc = v
        if isinstance(ws.cell(r, 4).value, (int, float)):
            mol = str(ws.cell(r, 2).value or '').strip()
            if mol in exclude:
                continue
            if nuc_filter is None or last_nuc == nuc_filter:
                rows.append(r)
    ref = np.array([ws.cell(r, 4).value for r in rows], dtype=float)
    last_f = None; cols = []
    for c in range(7, ws.max_column + 1):
        h = ws.cell(1, c).value
        if h is not None:
            last_f = h
        cols.append((c, last_f, ws.cell(2, c).value))
    data = {}
    for c, f, sub in cols:
        if f is None:
            continue
        vals = []; ok = True
        for r in rows:
            v = ws.cell(r, c).value
            if not isinstance(v, (int, float)):
                ok = False; break
            vals.append(v)
        if not ok or (f, sub) in data:
            continue
        data[(f, sub)] = np.array(vals, dtype=float)
    return ref, data

def load_ns372(path):
    """Load the NS372 (Kaupp) shielding matrix, sorted by family, for each of the 8 nuclei."""
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for nuc, sheet, filt, excl in NUCLEI:
        ref, data = read_nucleus(wb, sheet, filt, excl)
        used = [m for m in METHODS if (m[0], m[1]) in data]
        cols = [data[(m[0], m[1])] for m in used]
        labels = [m[4] for m in used] + ['CCSD(T)']
        fams = [m[3] for m in used] + ['ref']
        M = np.column_stack(cols + [ref])
        M, labels, fams = sort_by_family(M, labels, fams)
        out[nuc] = dict(M=M, methods=labels, ref=ref, families=fams)
    return out


# ----------------------------------------------------------------------------
# main-text Figure 2D (convergence to the electronic-structure limit)
# ----------------------------------------------------------------------------
# The 1H / pcSseg-3 slice of the Kaupp SI, as a raw-named table, plus the two
# per-method statistics the 3D panel plots (slope and scaled RMSE vs CCSD(T)).
# These reuse the NS372 reader above; nothing new is read from disk.
def load_ns372_1h_convergence(path):
    """1H CCSD(T) reference plus each benchmarked functional's shieldings, as columns.

    Tau policy differs from the SI S3 leveling analysis: Figure 2D uses the tauD kinetic-energy-density
    treatment for every tau-dependent functional, whereas load_ns372/METHODS uses tauC for meta-GGAs.
    So this selects tauD when present, otherwise the functional's single column, keyed by raw
    functional name ('Hartree-Fock' renamed 'HF' for the labels).
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ref, data = read_nucleus(wb, 'S6 - 1H Shieldings')
    names = {m[0] for m in METHODS}                      # the 44 benchmarked functionals
    variants = {}                                        # functional -> {tau: values}
    for (func, tau), vals in data.items():
        if func in names:
            variants.setdefault(func, {})[tau] = vals
    cols = {'CCSD(T)': ref}
    for func, by_tau in variants.items():
        vals = by_tau['τD'] if 'τD' in by_tau else next(iter(by_tau.values()))
        cols['HF' if func == 'Hartree-Fock' else func] = vals
    return pd.DataFrame(cols)

def convergence_statistics(cc_df, reference='CCSD(T)'):
    """Per method: the slope of method ~ a*CCSD(T) + b, and the scaled RMSE that Figure 2D plots.

    The prediction is rescaled back into reference units, y_scaled = (method - b) / a, and compared
    to CCSD(T); this equals the residual RMSE divided by the slope, so it differs from scaled_rmse,
    which the S3 bar plot leaves in the method's own units. These are the x/y coordinates of the
    Figure 2D points.
    """
    ref = cc_df[reference].to_numpy(float)
    rows = []
    for method in cc_df.columns:
        if method == reference:
            continue
        y = cc_df[method].to_numpy(float)
        slope, intercept = np.polyfit(ref, y, 1)
        y_scaled = (y - intercept) / slope
        rows.append({'Method': method, 'Slope': slope,
                     'RMSE (scaled)': _rmse(ref, y_scaled)})
    return pd.DataFrame(rows)


# ----------------------------------------------------------------------------
# delta22 loader (reads the released file; shieldings are decoded on read)
# ----------------------------------------------------------------------------
DELTA22_REF = 'DSD-PBEP86'

def load_delta22(path):
    """Load the delta22 shielding matrix, sorted by family, for 1H and 13C sites."""
    with h5py.File(path, 'r') as f:
        mn = [m.decode() for m in f['conventional_nmr_method_names'][:]]
        order = ['hf', 'blyp_d3bj', 'bp86_d3bj', 'b97d3_d3bj', 'tpsstpss_d3bj',
                 'b3lyp_d3bj', 'pbe0_d3bj', 'm062x_d3', 'wb97xd', 'wp04', 'wc04',
                 'B2PLYP', 'mPW2PLYP', 'B2GP_PLYP', 'dsd_pbep86', 'revdsd_pbep86',
                 'mp2', 'dlpno_mp2']
        disp = {'hf': 'HF', 'blyp_d3bj': 'BLYP-D3', 'bp86_d3bj': 'BP86-D3',
                'b97d3_d3bj': 'B97D3', 'tpsstpss_d3bj': 'TPSS-D3', 'b3lyp_d3bj': 'B3LYP-D3',
                'pbe0_d3bj': 'PBE0-D3', 'm062x_d3': 'M06-2X-D3', 'wb97xd': 'ωB97X-D',
                'wp04': 'WP04', 'wc04': 'WC04', 'B2PLYP': 'B2PLYP', 'mPW2PLYP': 'mPW2PLYP',
                'B2GP_PLYP': 'B2GP-PLYP', 'dsd_pbep86': 'DSD-PBEP86',
                'revdsd_pbep86': 'revDSD-PBEP86', 'mp2': 'MP2', 'dlpno_mp2': 'DLPNO-MP2'}
        fam = {'hf': 'WFT', 'blyp_d3bj': 'GGA', 'bp86_d3bj': 'GGA', 'b97d3_d3bj': 'GGA',
               'tpsstpss_d3bj': 'mGGA', 'b3lyp_d3bj': 'GH', 'pbe0_d3bj': 'GH',
               'm062x_d3': 'GH', 'wb97xd': 'RSH', 'wp04': 'GH', 'wc04': 'GH',
               'B2PLYP': 'DH', 'mPW2PLYP': 'DH', 'B2GP_PLYP': 'DH', 'dsd_pbep86': 'DH',
               'revdsd_pbep86': 'DH', 'mp2': 'WFT', 'dlpno_mp2': 'WFT'}
        gas = {}
        for i, m in enumerate(mn):
            func, basis, smodel, _ = m.split(',')
            if smodel == 'gas':
                gas.setdefault(func, {})[basis] = i
        sel, labels, families = [], [], []
        for func in order:
            b = gas[func]
            big = 'pcSseg3' if 'pcSseg3' in b else ('pcSseg2' if 'pcSseg2' in b else 'pcSseg1')
            sel.append(b[big]); labels.append(disp[func]); families.append(fam[func])
        GEOM = 1  # geometries axis: 0 = AIMNet2, 1 = PBE0/cc-pVTZ
        rows = {1: [], 6: []}
        for s in f['solutes']:
            an = f['solutes'][s]['atomic_numbers'][:]
            cs = _decode_fixed_point(
                f['solutes'][s]['stationary_and_pcm']['conventional_shieldings'][:, GEOM, :])
            for ai, z in enumerate(an):
                if z in (1, 6):
                    rows[int(z)].append(cs[sel, ai])
    out = {}
    for nuc, z in (('1H', 1), ('13C', 6)):
        M = np.array(rows[z])
        M, labs, fams = sort_by_family(M, list(labels), list(families))
        # use DSD-PBEP86 as the reference vector for delta22 so the scaled RMSE
        # bar plot mirrors NS372's CCSD(T) role
        ref_idx = labs.index(DELTA22_REF)
        out[nuc] = dict(M=M, methods=labs, ref=M[:, ref_idx], families=fams)
    return out
